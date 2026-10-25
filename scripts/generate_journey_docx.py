"""
Script to generate the comprehensive DecisionAI Project Journey Word Document (.docx).
Covers Phase 0 through current phase, detailing architecture, implementation,
problems encountered, and technical solutions.
Strict constraints: Zero emojis, zero em dashes.
"""

import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal padding for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def add_callout_box(doc, text, title="PROBLEM & SOLUTION HIGHLIGHT"):
    """Add a professional styled callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F7FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border styling
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2B6CB0"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"{title}: ")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(43, 108, 176)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(45, 55, 72)
    
    # Empty paragraph after table for spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def build_document():
    doc = docx.Document()
    
    # Page setup: Normal 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    # Base palette
    COLOR_PRIMARY = RGBColor(26, 54, 93)     # Deep Navy (#1A365D)
    COLOR_SECONDARY = RGBColor(43, 108, 176) # Accent Blue (#2B6CB0)
    COLOR_TEXT = RGBColor(45, 55, 72)        # Charcoal (#2D3748)
    COLOR_MUTED = RGBColor(113, 128, 150)    # Gray (#718096)
    
    # Title Section
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("DecisionAI: Complete Engineering Journey")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(26)
    run_title.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(18)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("From Project Inception (Phase 0) to Thought-Partner Meeting Intelligence")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = COLOR_SECONDARY
    
    # Metadata Box
    meta_tbl = doc.add_table(rows=1, cols=3)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    col_widths = [Inches(2.1), Inches(2.3), Inches(2.1)]
    meta_items = [
        ("Architecture", "FastAPI + MV3 + ChromaDB"),
        ("Coverage", "Phase 0 through Phase 8"),
        ("Status", "Production Grade (61/61 Tests Passing)")
    ]
    for i, (label, val) in enumerate(meta_items):
        cell = meta_tbl.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "EDF2F7")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        lbl_r = p.add_run(f"{label}\n")
        lbl_r.font.name = "Calibri"
        lbl_r.font.size = Pt(9)
        lbl_r.bold = True
        lbl_r.font.color.rgb = COLOR_MUTED
        val_r = p.add_run(val)
        val_r.font.name = "Calibri"
        val_r.font.size = Pt(9.5)
        val_r.bold = True
        val_r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper for adding headings
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(18)
        r.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(14)
        r.bold = True
        r.font.color.rgb = COLOR_SECONDARY
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11.5)
        r.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return h

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            br = p.add_run(bold_prefix)
            br.bold = True
            br.font.name = "Calibri"
            br.font.size = Pt(11)
            br.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            br = p.add_run(bold_prefix)
            br.bold = True
            br.font.name = "Calibri"
            br.font.size = Pt(11)
            br.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_TEXT
        return p

    # --- EXECUTIVE SUMMARY ---
    add_h1("1. Executive Summary & Project Vision")
    add_p(
        "DecisionAI was conceived to eliminate the limitations of traditional meeting recording bots. "
        "Conventional meeting bots (such as Otter, Fireflies, or Zoom bots) require external calendar invitations, "
        "appear as intrusive participant accounts in video calls, and rely on cloud audio streams that often drop "
        "nuanced multi-lingual dialogue. In bilingual and multilingual engineering teams, participants switch "
        "effortlessly between English, Hindi, and Hinglish. Traditional speech engines and storage databases fail "
        "to transcribe code-switching accurately, omit contextual speaker attributions, and behave as passive text dumps."
    )
    add_p(
        "The core mission of DecisionAI was defined by four principles:",
        "Four Architectural Pillars: "
    )
    add_bullet("Direct Browser Audio Capture: Capture participant tab audio and microphone audio directly from Google Meet without meeting bots.", "1. Non-Intrusive Ingestion: ")
    add_bullet("Hardware-Accelerated Dual Worklets: Resample audio in real-time to 16 kHz mono 16-bit PCM for clean streaming.", "2. Signal Precision: ")
    add_bullet("Pluggable STT with Code-Switching: Accurately process English, Devanagari Hindi, and Hinglish code-switching.", "3. Linguistic Inclusivity: ")
    add_bullet("Intelligent Thought Partner: Function as an active assistant that extracts decisions, tracks action items, and generates valuable technical recommendations rather than serving as a passive archive.", "4. Cognitive Value: ")

    # --- PHASE 0 ---
    add_h1("2. Phase 0: Foundation, Architecture & Environment Setup")
    add_p(
        "Phase 0 established the core repository structure, configuration management, and foundational design standards. "
        "The architecture is anchored on FastAPI for asynchronous high-concurrency request routing, SQLite with aiosqlite "
        "for zero-infrastructure prototype persistence, and Pydantic Settings for centralized environment configuration."
    )
    add_h2("What Was Built in Phase 0")
    add_bullet("Centralized configuration system in backend/config.py loading environment variables via pydantic-settings.", "Settings & Environment: ")
    add_bullet("FastAPI application instance with CORS middleware, asynchronous lifecycle handlers, and error handlers.", "FastAPI Core Application: ")
    add_bullet("Asynchronous database engine factory with SQLite connection pooling and write-ahead logging (WAL).", "Database Engine Foundation: ")
    add_bullet("Strict system-wide design constraints: zero third-party UI framework dependencies, WCAG AA contrast compliance, and no em-dash or emoji artifacts.", "Development Constraints: ")

    add_callout_box(
        doc,
        "During initial setup, conflicting environment variable loading between local .env files and process variables caused unpredictable database paths in testing. "
        "Solution: Standardized on pydantic-settings with explicit env_file='.env' configuration, priority override hierarchies, and test-scoped in-memory database URLs.",
        "PHASE 0 CHALLENGE & RESOLUTION"
    )

    # --- PHASE 1 ---
    add_h1("3. Phase 1: Relational Data Models & Asynchronous Persistence")
    add_p(
        "Phase 1 implemented the complete relational schema to represent the lifecycle of meetings, participants, dialogue segments, "
        "and extracted business intelligence. SQLAlchemy 2.0 declarative models were created with strict relationship constraints."
    )
    add_h2("Core Models Built")
    add_bullet("Represents registered users, storing email, OAuth identifiers, hashed credentials, and account creation dates.", "User Model: ")
    add_bullet("Tracks recording sessions with status transitions (CREATED -> RECORDING -> PROCESSING -> READY -> FAILED), duration metrics, and audio file paths.", "Meeting Model: ")
    add_bullet("Stores individual speech turns with speaker identity (Mic vs Tab), start/end millisecond offsets, raw text, normalized text, and language tags.", "TranscriptSegment Model: ")
    add_bullet("Stores business decisions, action items, assignees, deadlines, and review status (accepted, rejected, pending).", "Decision Model: ")
    add_bullet("Tracks security-critical events (recording started, audio downloaded, data deleted) with IP addresses and timestamps.", "AuditLog Model: ")

    add_callout_box(
        doc,
        "Asynchronous session management in SQLAlchemy 2.0 frequently caused greenlet spawning errors when accessing lazy-loaded relationships during serialization. "
        "Solution: Enforced explicit selectinload queries and strict Pydantic DTO schemas across all API responses, completely avoiding lazy-load boundary leaks.",
        "PHASE 1 CHALLENGE & RESOLUTION"
    )

    # --- PHASE 2 ---
    add_h1("4. Phase 2: Audio File Upload & Storage Management")
    add_p(
        "Phase 2 introduced manual audio file ingestion to support legacy recordings, offline testing, and audio validation. "
        "A multipart file upload endpoint (/api/meetings/upload) was built with rigorous security checks."
    )
    add_h2("Features Implemented")
    add_bullet("Validation of audio MIME types (WAV, MP3, M4A, WebM) and rejection of empty or non-audio binary payloads.", "File Validation: ")
    add_bullet("Enforcement of 500 MB upload limits to prevent memory exhaustion and disk quota denial-of-service attacks.", "Storage Quotas: ")
    add_bullet("Automated duration extraction using soundfile and wave libraries to establish meeting timelines.", "Audio Analysis: ")

    add_callout_box(
        doc,
        "Different operating systems and browser clients sent inconsistent MIME types for audio files (e.g. audio/x-wav vs audio/wav, audio/wave). "
        "Solution: Implemented magic-byte header inspection combined with file extension whitelisting, ensuring reliable validation across all platforms.",
        "PHASE 2 CHALLENGE & RESOLUTION"
    )

    # --- PHASE 3 ---
    add_h1("5. Phase 3: Manifest V3 Chrome Extension & Dual-Channel Live Capture")
    add_p(
        "Phase 3 represented the most critical hardware and browser integration hurdle of the project. "
        "Capturing dual-channel call audio directly from Google Meet without an external bot required navigating Chrome Manifest V3 restrictions."
    )
    add_h2("Chrome Extension Architecture")
    add_bullet("Executes in background service worker upon user initiation, generating a cryptographically restricted streamId.", "chrome.tabCapture API: ")
    add_bullet("Because Manifest V3 service workers lack DOM and Web Audio APIs, an offscreen document is spawned with USER_MEDIA reason.", "Offscreen Document: ")
    add_bullet("Dual AudioContext instances run dedicated AudioWorklets (PCMRecorderProcessor) for Tab and Mic channels, downmixing to mono and converting Float32 to 16 kHz Int16 PCM.", "Hardware AudioWorklet: ")
    add_bullet("Continuous binary streaming over WebSocket framed with 4-byte ASCII headers (TAB: and MIC:), backed by an in-memory reconnect buffer queue.", "Framed WebSocket Streaming: ")

    add_callout_box(
        doc,
        "Capturing tab audio via chrome.tabCapture automatically muted the call in Chrome! The user could record call participants, but could no longer hear them in their own headphones. "
        "Solution: In offscreen.js, the captured participant MediaStream was split and connected to audioCtx.destination. This maintained clear passthrough audio to the user's headphones while simultaneously piping samples to the recording worklet.",
        "PHASE 3 MAJOR PROBLEM & SOLUTION: TAB AUDIO MUTING"
    )

    add_callout_box(
        doc,
        "During live streaming, if the user had not yet authenticated in the web application, the WebSocket connection disconnected immediately with HTTP 403, causing the recording to freeze and save 0.0 seconds of audio. "
        "Solution: Implemented the /auth/extension-token endpoint. The popup authenticates the user, obtains a secure stream token, passes it to the offscreen document, and injects it into the WebSocket handshake query parameters.",
        "PHASE 3 MAJOR PROBLEM & SOLUTION: WEBSOCKET AUTHENTICATION"
    )

    # --- PHASE 4 ---
    add_h1("6. Phase 4: Pluggable Multi-Provider STT & Multilingual Normalization")
    add_p(
        "Phase 4 delivered a multi-provider Speech-to-Text architecture capable of accurately transcribing conversations across English, "
        "Devanagari Hindi, and colloquial Hinglish code-switching."
    )
    add_h2("Provider Ecosystem")
    add_bullet("Local CTranslate2 inference with int8/float16 quantization. Completely offline, zero cost, and low latency.", "faster-whisper: ")
    add_bullet("Multimodal audio understanding via Google GenAI REST API. Handles nuanced conversational code-switching and structured JSON timestamps.", "Google Gemini: ")
    add_bullet("Specialist Indian language engine (saaras:v4) operating in native codemix mode with glossary prompt hints.", "Sarvam AI: ")
    add_bullet("Cloud enterprise platform with nova-3 and keyword boosting for high-volume English dialogue.", "Deepgram: ")

    add_h2("Channel Merge and Normalization Pipeline")
    add_p(
        "The background processor chunks audio into 30-second windows with 0.5-second overlap. Both Tab and Mic audio are transcribed concurrently. "
        "Recognized speech segments from both channels are chronologically sorted and merged by start_ms, assigning speaker tags (You vs Participant). "
        "A multilingual normalizer applies Unicode NFC normalization, standardizes Hindi nuktas and dandas, and resolves phonetic Hinglish variations."
    )

    add_callout_box(
        doc,
        "When a user was alone in a Google Meet call explaining a project, the participant Tab channel recorded pure digital silence (RMS amplitude = 0.00). "
        "Gemini STT previously hallucinated random 1878 historical text (Vernacular Press Act) on pure silence! "
        "Solution: Implemented is_silent_wav pre-flight energy gating. The pipeline computes Root-Mean-Square (RMS) amplitude across the audio buffer. If pure digital silence is detected, transcription is skipped with zero hallucination. Furthermore, Gemini prompts were updated with strict silence suppression instructions.",
        "PHASE 4 MAJOR PROBLEM & SOLUTION: SILENCE HALLUCINATIONS"
    )

    # --- PHASE 5 ---
    add_h1("7. Phase 5: Authentication, User Scoping & Enterprise Data Governance")
    add_p(
        "Phase 5 established complete multi-tenant security, user scoping, and privacy governance to meet enterprise compliance standards."
    )
    add_h2("Security Architecture")
    add_bullet("Authenticated using official Google Client ID and Secret, setting httpOnly=True, samesite='lax' session cookies.", "Real Google OAuth 2.0: ")
    add_bullet("Every database query, meeting view, and ChromaDB vector search strictly enforces a WHERE user_id = :user_id constraint.", "Multi-Tenant Isolation: ")
    add_bullet("In-memory sliding window rate limiter without external Redis dependencies, capping requests at 60/min per IP/token.", "Sliding-Window Rate Limiting: ")
    add_bullet("Validates meeting IDs, string lengths, and payload boundaries against SQL injection and buffer overflow exploits.", "Input Validation: ")
    add_bullet("Single-click complete data wipeout endpoint (/api/user/delete-all) purging all meetings, transcripts, decisions, and audio files.", "Privacy Compliance: ")

    add_callout_box(
        doc,
        "CSRF protection middleware on the FastAPI backend intercepted and blocked legitimate POST requests dispatched by the Chrome extension (such as /meetings/finish and WebSocket upgrades). "
        "Solution: Configured selective CSRF verification that validates browser cookie-based web dashboard requests while exempting extension endpoints authenticated via Bearer extension tokens.",
        "PHASE 5 CHALLENGE & RESOLUTION: CSRF EXEMPTION"
    )

    # --- PHASE 6 ---
    add_h1("8. Phase 6: Post-Processing Intelligence & Structured Decisions Extraction")
    add_p(
        "Phase 6 transformed raw multi-speaker transcripts into structured, actionable business intelligence. Following transcription, "
        "an asynchronous worker executes Gemini LLM post-processing."
    )
    add_h2("Extracted Intelligence Components")
    add_bullet("Technical glossary hints (e.g. DecisionAI, ChromaDB, AudioWorklet) correct phonetic misrecognitions while preserving raw_text.", "Glossary Cleanup: ")
    add_bullet("A concise synthesis of meeting discussions, participant positions, and core conclusions.", "Executive Summary: ")
    add_bullet("Detection of agreed decisions, assigned tasks, deliverable deadlines, and review status (accepted, rejected, pending).", "Decisions & Action Items: ")
    add_bullet("Structured parsing validated against strict Pydantic schemas, with deterministic offline fallback when API limits occur.", "Schema Validation: ")

    add_callout_box(
        doc,
        "During live recording finalization, the UI popup previously froze at 86% processing and never redirected to the meeting viewer. "
        "Solution: Investigated polling logic in popup.js. The frontend was polling the generic meetings list rather than checking the specific meeting ID status, causing state lockups. Updated popup polling to target /api/meetings/{id} directly and automatically redirect the browser tab upon status reaching 'ready'.",
        "PHASE 6 CHALLENGE & RESOLUTION: 86% PROGRESS FREEZE"
    )

    # --- PHASE 7 ---
    add_h1("9. Phase 7: Responsive Vanilla Frontend & Bento Grid Dashboard")
    add_p(
        "Phase 7 focused on the presentation layer, building an ultra-fast, zero-framework web dashboard (/web) adhering to strict design standards."
    )
    add_h2("Frontend Design System")
    add_bullet("Clean solid borders, subtle contrast hierarchy, and complete elimination of blur and heavy shadows.", "Flat Solid Surface: ")
    add_bullet("Modular CSS Grid cards organizing recent meetings, quick stats, active decisions, and Ask AI entry points.", "Bento Grid Dashboard: ")
    add_bullet("All color tokens verified for WCAG AA compliance (contrast ratios exceeding 5.0:1 for primary and secondary copy).", "Accessibility (WCAG AA): ")
    add_bullet("Integrated synchronized dual-channel audio player allowing playback seeking, volume adjustments, and transcript highlighting.", "Meeting Viewer: ")

    add_callout_box(
        doc,
        "Technical jargon (such as 'Int16 PCM Worklet' and 'RRF k=60 Fusion') was initially displayed to end users in the interface, causing confusion. "
        "Solution: Undertook a comprehensive copy refactor across the entire web app and extension, replacing internal engineering terms with intuitive executive language (e.g. 'Processing audio', 'Capturing call', 'Ask AI').",
        "PHASE 7 REFACTOR: USER-FRIENDLY TERMINOLOGY"
    )

    # --- PHASE 8 ---
    add_h1("10. Phase 8: Hybrid RAG, Thought-Partner Advisory AI & Resilience")
    add_p(
        "Phase 8 established the conversational intelligence engine of DecisionAI, elevating the system from a passive transcription database "
        "into an intelligent co-pilot and thought partner."
    )
    add_h2("RAG Pipeline Architecture")
    add_bullet("Transcripts grouped into 300 to 500 token chunks along speaker turn boundaries, preserving conversational context.", "Speaker-Turn Chunking: ")
    add_bullet("ChromaDB vector store backed by paraphrase-multilingual-MiniLM-L12-v2 with mandatory multi-tenant user_id filtering.", "Vector Retrieval: ")
    add_bullet("In-memory BM25 index supporting English, Hindi (Devanagari), and romanized Hinglish tokens.", "Sparse Lexical Search: ")
    add_bullet("Dense and sparse scores are fused fairly using Reciprocal Rank Fusion with constant k=60.", "Reciprocal Rank Fusion (RRF): ")
    add_bullet("Answers cite exact meeting titles, dates, and MM:SS timestamps with clickable interactive chips that seek audio playback directly.", "Grounded Citations: ")

    add_callout_box(
        doc,
        "When the user asked consultative or advisory questions (such as 'do u think we can add anything else valuable to it'), the AI refused to answer with: 'No relevant meeting discussions found matching...'. "
        "The system was behaving like a rigid database lookup tool rather than an intelligent meeting assistant. "
        "Solution: Re-architected the prompt and intent pipeline. Added is_advisory_query classification, separating factual recall queries (which stay faithful to recorded facts) from advisory queries (which synthesize the project context and apply AI intellect and reasoning to propose concrete architectural recommendations). Furthermore, updated offline fallback synthesis to generate domain-tailored suggestions rather than quitting.",
        "PHASE 8 MAJOR PROBLEM & SOLUTION: THOUGHT-PARTNER ADVISORY MODE"
    )

    add_callout_box(
        doc,
        "Google Gemini free-tier experienced intermittent HTTP 429 quota limits and HTTP 503 high-demand errors, causing RAG queries to fail unexpectedly. "
        "Solution: Implemented multi-model candidate cascading across gemini-3.5-flash-lite, gemini-flash-lite-latest, gemini-3.1-flash-lite, and gemini-3.5-flash with exponential backoff retries and dynamic temperature adjustment (0.4 for brainstorming, 0.2 for factual recall).",
        "PHASE 8 MAJOR PROBLEM & SOLUTION: MULTI-MODEL QUOTA RESILIENCE"
    )

    # --- SUMMARY TABLE ---
    add_h1("11. Summary of Project Journey & Milestones")
    
    sum_tbl = doc.add_table(rows=1, cols=4)
    sum_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sum_tbl.autofit = False
    headers = ["Phase", "Milestone / Scope", "Key Challenge Encountered", "Technical Solution"]
    widths = [Inches(1.0), Inches(1.8), Inches(1.8), Inches(1.9)]
    
    for i, h in enumerate(headers):
        cell = sum_tbl.cell(0, i)
        cell.width = widths[i]
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    milestones = [
        ("Phase 0", "Project Skeleton & Config", "Inconsistent config paths across OS environments", "Pydantic Settings with centralized single source of truth"),
        ("Phase 1", "Database Schema & CRUD", "SQLAlchemy 2.0 greenlet spawning on lazy relationships", "Explicit selectinload and dedicated Pydantic response DTOs"),
        ("Phase 2", "Audio Upload & Storage", "MIME type variations across operating systems", "Magic-byte binary inspection and size limit validation"),
        ("Phase 3", "Chrome MV3 Live Capture", "Tab capture automatically muted browser call audio", "Split stream to audioCtx.destination in offscreen document"),
        ("Phase 4", "Pluggable STT Pipeline", "Gemini STT hallucinated historical text on silent calls", "Root-Mean-Square (RMS) pre-flight digital silence gating"),
        ("Phase 5", "Auth & Data Governance", "CSRF middleware blocked Chrome extension requests", "Exempted Bearer token authenticated extension endpoints"),
        ("Phase 6", "Intelligence Extraction", "Popup UI froze at 86% during meeting processing", "Targeted meeting ID polling and automated browser redirect"),
        ("Phase 7", "Vanilla Web Dashboard", "Technical jargon confused executive end users", "Comprehensive copy refactoring across UI and extension"),
        ("Phase 8", "Thought-Partner RAG", "Assistant refused consultative questions as 'not found'", "Added is_advisory_query, thought-partner prompt, and 429 backoff")
    ]

    for phase, scope, chal, sol in milestones:
        row = sum_tbl.add_row()
        for idx, text in enumerate([phase, scope, chal, sol]):
            cell = row.cells[idx]
            cell.width = widths[idx]
            set_cell_background(cell, "F7FAFC" if len(sum_tbl.rows) % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_TEXT

    # Output path
    output_path = Path("c:/Users/sanid/DecisionIQ/DecisionAI/DecisionAI_Project_Journey.docx")
    doc.save(str(output_path))
    print(f"Document successfully saved to {output_path}")

if __name__ == "__main__":
    build_document()
